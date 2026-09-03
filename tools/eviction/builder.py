"""Builds the NC Summary Ejectment court-prep documents (python-docx).

  build_quickref(case, out)  -> 1-2 page courtroom quick-reference card
  build_packet(case, out)    -> full multi-part prep binder

Ported from the nc-eviction-court-packet skill; the NC legal scaffolding
(pay-and-stay G.S. 42-33, money-judgment service rule, proration, defenses)
is preserved.
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from . import casecalc

NAVY = "1F3864"; SLATE = "44546A"; LIGHT = "D9E1F2"; LIGHTER = "EAF0FA"
WARN = "FCE4D6"; WARNB = "C55A11"; GREEN = "E2EFDA"; GREENB = "70AD47"; RULE = "8EAADB"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# ---------- low-level helpers ----------
def _shade(el, fill):
    pr = el.get_or_add_pPr() if el.tag.endswith('}p') else el
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill)
    pr.append(sh)

def shade_p(p, fill):
    _shade(p._p, fill)

def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill)
    tcPr.append(sh)

def cell_borders(cell, **sides):
    """sides: left/right/top/bottom = (size_eighths, color) or None"""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for name in ("top", "left", "bottom", "right"):
        spec = sides.get(name)
        e = OxmlElement('w:' + name)
        if spec:
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(spec[0]))
            e.set(qn('w:color'), spec[1])
        else:
            e.set(qn('w:val'), 'nil')
        borders.append(e)
    tcPr.append(borders)

def no_table_borders(table):
    tbl = table._tbl
    pr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for n in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement('w:' + n); e.set(qn('w:val'), 'nil'); borders.append(e)
    pr.append(borders)

def sp(p, before=0, after=4):
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)

def run(p, text, size=11, bold=False, italic=False, color=None):
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color:
        r.font.color.rgb = color if isinstance(color, RGBColor) else RGBColor.from_string(color)
    return r

def rich(p, parts, size=11):
    """parts: list of (text, {bold/italic/color}) or plain strings"""
    for part in parts:
        if isinstance(part, str):
            run(p, part, size=size)
        else:
            txt, o = part
            run(p, txt, size=o.get("size", size), bold=o.get("bold", False),
                italic=o.get("italic", False), color=o.get("color"))
    return p

def _footer(section, text):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, text + " — page ", size=7, color="808080")
    for instr in ("PAGE",):
        fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), instr)
        r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
        szv = OxmlElement('w:sz'); szv.set(qn('w:val'), "14"); rPr.append(szv)
        c = OxmlElement('w:color'); c.set(qn('w:val'), "808080"); rPr.append(c)
        r.append(rPr); fld.append(r); p._p.append(fld)
    run(p, " of ", size=7, color="808080")
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), "NUMPAGES")
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    szv = OxmlElement('w:sz'); szv.set(qn('w:val'), "14"); rPr.append(szv)
    c = OxmlElement('w:color'); c.set(qn('w:val'), "808080"); rPr.append(c)
    r.append(rPr); fld.append(r); p._p.append(fld)

# ---------- block builders ----------
def bar(doc, text, fill=NAVY, size=10.5):
    p = doc.add_paragraph(); sp(p, 6, 3); shade_p(p, fill)
    run(p, "  " + text.upper(), size=size, bold=True, color=WHITE)
    return p

def h1(doc, text):
    p = doc.add_paragraph(); sp(p, 16, 6)
    run(p, text, size=14, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom'); b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), '12'); b.set(qn('w:color'), RULE); pbdr.append(b); pPr.append(pbdr)
    return p

def h2(doc, text):
    p = doc.add_paragraph(); sp(p, 10, 4)
    run(p, text, size=12, bold=True, color=SLATE)
    return p

def para(doc, parts, size=11, after=6, italic=False, color=None):
    p = doc.add_paragraph(); sp(p, 0, after)
    if isinstance(parts, str):
        run(p, parts, size=size, italic=italic, color=color)
    else:
        rich(p, parts, size=size)
    return p

def bullet(doc, parts, size=11):
    p = doc.add_paragraph(style="List Bullet"); sp(p, 0, 3)
    if isinstance(parts, str):
        run(p, parts, size=size)
    else:
        rich(p, parts, size=size)
    return p

def numbered(doc, parts, size=11):
    p = doc.add_paragraph(style="List Number"); sp(p, 0, 4)
    if isinstance(parts, str):
        run(p, parts, size=size)
    else:
        rich(p, parts, size=size)
    return p

def callout(doc, title, lines, fill=WARN, border=WARNB, width=7.1):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]; cell.width = Inches(width)
    shade_cell(cell, fill)
    cell_borders(cell, top=(12, border), bottom=(12, border), left=(24, border), right=(12, border))
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
    if title:
        p = cell.add_paragraph(); sp(p, 2, 3)
        run(p, title, size=10.5, bold=True, color="833C00" if fill == WARN else SLATE)
    for line in lines:
        p = cell.add_paragraph(); sp(p, 0, 3)
        if isinstance(line, str):
            run(p, line, size=10)
        else:
            rich(p, line, size=10)
    return t

def fact_table(doc, rows, label_w=2.3, width=7.1):
    t = doc.add_table(rows=len(rows), cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_table_borders(t)
    for i, (label, val) in enumerate(rows):
        c0, c1 = t.rows[i].cells
        c0.width = Inches(label_w); c1.width = Inches(width - label_w)
        fill = "FFFFFF" if i % 2 else LIGHTER
        shade_cell(c0, fill); shade_cell(c1, fill)
        cell_borders(c0, bottom=(2, "C9D3EA")); cell_borders(c1, bottom=(2, "C9D3EA"))
        p0 = c0.paragraphs[0]; sp(p0, 1, 1); run(p0, str(label), size=9.5, bold=True, color=SLATE)
        p1 = c1.paragraphs[0]; sp(p1, 1, 1)
        if isinstance(val, str):
            run(p1, val, size=9.5)
        else:
            rich(p1, val, size=9.5)
    return t

def grid_table(doc, header, rows, widths, width=7.1):
    t = doc.add_table(rows=1 + len(rows), cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(header):
        c = t.rows[0].cells[i]; c.width = Inches(widths[i])
        shade_cell(c, NAVY)
        p = c.paragraphs[0]; sp(p, 1, 1); run(p, htxt, size=9, bold=True, color=WHITE)
    for ri, r in enumerate(rows):
        fill = "FFFFFF" if ri % 2 else LIGHTER
        for ci, val in enumerate(r):
            c = t.rows[ri + 1].cells[ci]; c.width = Inches(widths[ci])
            shade_cell(c, fill)
            p = c.paragraphs[0]; sp(p, 1, 1)
            if isinstance(val, str):
                run(p, val, size=9)
            else:
                rich(p, val, size=9)
    return t

def _setup(doc, margins=0.6):
    for s in doc.sections:
        s.top_margin = Inches(margins); s.bottom_margin = Inches(margins)
        s.left_margin = Inches(margins + 0.05); s.right_margin = Inches(margins + 0.05)
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    return doc

def _landlord(land):
    """'the Smith Family Trust' vs 'Bartola Lisbon' — controlled by landlord.article."""
    name = land.get("name", "")
    art = land.get("article")
    if art is None:
        # infer: entities/trusts take "the"
        low = name.lower()
        art = "the " if any(w in low for w in
              ("trust", "llc", "l.l.c", "inc", "corp", "company", "properties", "partners", "lp", "estate")) else ""
    return "%s%s" % (art, name)

def _pretty(v):
    """Render a date value as 'August 1, 2026' when parseable, else as given."""
    d = casecalc._d(v)
    if not d:
        return str(v or "")
    return "%s %d, %d" % (d.strftime("%B"), d.day, d.year)

def _short_date(v):
    """Short form for inline prose: 'August 6' (drops the current year)."""
    d = casecalc._d(v)
    if not d:
        return str(v or "").replace(", 2026", "")
    return "%s %d" % (d.strftime("%B"), d.day)

# ============================================================
#  QUICK REFERENCE CARD
# ============================================================
def build_quickref(case, out_path):
    d = casecalc.compute(case)
    money = casecalc.money
    firm = case.get("firm", {}); land = case.get("landlord", {}); ten = case.get("tenant", {})
    prop = case.get("property", {}); court = case.get("court", {}); rent = case.get("rent", {})
    claim = case.get("claim", {}); demand = case.get("demand", {}); filing = case.get("filing", {})

    doc = Document(); _setup(doc, 0.5)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 0, 1)
    run(p, "COURTROOM QUICK REFERENCE", size=14, bold=True, color=NAVY)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 0, 1)
    run(p, "%s  v.  %s" % (land.get("name", "Landlord"), ten.get("name", "Tenant")), size=10, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 0, 4)
    run(p, "%s • File %s • %s %s • %s, %s" % (
        court.get("countyCourt", ""), court.get("fileNo", ""),
        str(court.get("hearingDate", "")).split(", ", 1)[-1], court.get("hearingTime", ""),
        court.get("courthouseShort", ""), court.get("room", "")), size=8.5, color=SLATE)

    bar(doc, "Say this when your case is called")
    callout(doc, None, [
        [('"My name is ______________________ and I am acting as the agent for the landlord of this property. The landlord is ', {}),
         (_landlord(land), {}), (", and my company, %s, is the management company." % firm.get("name", ""), {})],
        [("The tenant, %s, rents %s. The rent is %s a month, due on %s. " % (
            ten.get("name", ""), prop.get("address", ""), d["FULL"], rent.get("dueDayText", "the 1st")), {}),
         ("They did not pay the %s, which was due %s, and it is still unpaid. We sent a written demand on %s and filed on %s." % (
             claim.get("rentMonthLabel", "rent"), _pretty(claim.get("rentDueDate")),
             _short_date(demand.get("sentDate")), _short_date(filing.get("filedDate"))), {})],
        [("We are asking for possession of the property and a judgment for the unpaid rent plus court costs. "
          "I have the lease, the ledger, and the demand notice here — may I enter them into evidence?\"", {})],
    ], LIGHT, RULE)

    bar(doc, "The ask  —  say both out loud")
    ask_lines = [
        [("1.  POSSESSION ", {"bold": True}), ("of ", {}), (prop.get("address", ""), {})],
        [("2.  MONEY JUDGMENT ", {"bold": True}), ("for rent through the court date = ", {}), (d["PRO"], {"bold": True}),
         ("   (%d days × %s accrued = %s%s)" % (
             d["prorated_days"], d["D"], d["ACCRUED"],
             "; less %s already paid toward this rent" % money(d["current_paid"]) if d["current_paid"] else ""), {})],
        [("     + COURT COSTS ", {"bold": True}), (d["COSTS"], {"bold": True}), ("   (%s)" % d["costs_breakdown"], {})],
        [("     TOTAL MONEY ASK ", {"bold": True}), (d["TOTAL"], {"bold": True})],
        [("     If the magistrate awards the whole month instead: ", {}), (d["MONTHASK"], {"bold": True}),
         (" rent (%s charged%s)" % (d["CHARGED"], (" − %s paid" % d["PAID"]) if d["current_paid"] else ""), {})],
    ]
    if d["prior_rent_unpaid"]:
        ask_lines.insert(2, [("     includes ", {}), (money(d["prior_rent_unpaid"]), {"bold": True}),
                             (" of unpaid rent from a prior month", {})])
    callout(doc, None, ask_lines, GREEN, GREENB)

    if d["has_payments"]:
        pay_lines = [[("A payment was received and applied to the oldest outstanding charge first:", {"bold": True})]]
        for a in d["allocations"]:
            pay_lines.append([("   %s  %s  →  %s" % (a["payment_date"], money(a["amount"]), a["applied_to"]), {})])
        if d["unapplied"]:
            pay_lines.append([("   %s not applied to any charge — verify with the office." % money(d["unapplied"]), {"bold": True})])
        pay_lines.append([("A partial payment does NOT stop the case. ", {"bold": True}),
                          ("Disclose it to the magistrate, and confirm with counsel how it affects the claim "
                           "(see the lease's partial-rent paragraph).", {})])
        bar(doc, "Partial payment received — read this")
        callout(doc, None, pay_lines, WARN, WARNB)

    bar(doc, "Your proof — hand up as you go")
    bullet(doc, [("Lease ", {"bold": True}), ("(%s) — the tenancy, rent, due date, breach terms" % case.get("tenancy", {}).get("leaseForm", ""), {})], 10)
    bullet(doc, [("Ledger ", {"bold": True}), ("— rent charged, payments applied, balance owed", {})], 10)
    bullet(doc, [("Demand notice ", {"bold": True}), ("(%s) — written demand for possession" % _short_date(demand.get("sentDate")), {})], 10)
    bullet(doc, [("File-stamped complaint ", {"bold": True}), ("(filed %s)" % _short_date(filing.get("filedDate")), {})], 10)

    bar(doc, "If they ask for more time  /  offer to pay")
    callout(doc, None, [
        [("Default: press for judgment today. ", {"bold": True}),
         ("Don't agree to a continuance or payment plan for a promise. Winning now isn't harsh — "
          "they still have 10 days before any lockout to pay in full.", {})],
        [("Pays in full before the ruling (rent + costs)? ", {"bold": True}),
         ("That ends the case by law (G.S. 42-33) — accept it, get certified funds if you can, and call the office.", {})],
        [("Partial payment or a promise? ", {"bold": True}),
         ("Does NOT stop it — proceed. Don't cut any side deal on your own; step out and call the office.", {})],
    ], WARN, WARNB)

    bar(doc, "Before you rely on the money judgment")
    para(doc, [("Money judgment needs personal service OR their appearance. ", {"bold": True}),
               ("Check the sheriff's return. If served only by posting and they don't show, the court can grant "
                "possession only — not the money. If they're in the room, the money judgment is available; make the full ask.", {})], size=10)

    bar(doc, "Logistics")
    bullet(doc, [("Leave your phone in the car ", {"bold": True}), ("— phones aren't allowed inside the courthouse. Bring everything on paper.", {})], 10)
    bullet(doc, [("Arrive ~%s " % court.get("arriveTime", ""), {"bold": True}),
                 ("for the %s calendar; allow time for parking + security. %s." % (court.get("hearingTime", ""), court.get("room", "")), {})], 10)
    bullet(doc, [("Bring 3 copies ", {"bold": True}), ("of each exhibit (you, magistrate, tenant). Have the management agreement in case authority is questioned.", {})], 10)
    bullet(doc, [("After a win: ", {"bold": True}), ("10-day appeal window (rolls to next business day if it lands on a weekend/holiday); no self-help lockout; report the result to the office.", {})], 10)

    bar(doc, "Numbers & dates")
    rows = [
        ("Monthly rent", "%s, due %s   •   Daily rent %s" % (d["FULL"], rent.get("dueDayText", "the 1st"), d["D"])),
        ("Ask for (rent)", [(d["PRO"], {"bold": True}), ("  through court date", {})]),
        ("Ask for (costs)", d["COSTS"]),
        ("Total money ask", [(d["TOTAL"], {"bold": True})]),
    ]
    if d["fees_outstanding"]:
        rows.append(("Fees NOT in the ask", "%s (late/add-on fees)" % money(d["fees_outstanding"])))
    rows += [
        ("Rent due / demand / filed", "%s  /  %s  /  %s" % (
            _pretty(claim.get("rentDueDate")), _short_date(demand.get("sentDate")), _short_date(filing.get("filedDate")))),
        ("Office", "%s — %s" % (firm.get("name", ""), firm.get("phone", ""))),
    ]
    fact_table(doc, rows, 2.0)
    p = doc.add_paragraph(); sp(p, 6, 0)
    run(p, "Internal prep aid — not legal advice. Full details in the Court-Prep Packet.", size=7.5, italic=True, color="808080")
    _footer(doc.sections[0], "%s — Quick Reference — File %s" % (
        str(ten.get("name", "")).split(" ")[-1], court.get("fileNo", "")))
    doc.save(out_path)
    return out_path

# ============================================================
#  FULL PREP BINDER
# ============================================================
def build_packet(case, out_path):
    d = casecalc.compute(case)
    money = casecalc.money
    firm = case.get("firm", {}); land = case.get("landlord", {}); ten = case.get("tenant", {})
    prop = case.get("property", {}); court = case.get("court", {}); rent = case.get("rent", {})
    claim = case.get("claim", {}); demand = case.get("demand", {}); filing = case.get("filing", {})
    tenancy = case.get("tenancy", {}); ledger = case.get("ledger", {})
    tname = ten.get("name", "Tenant"); tlast = str(tname).split(" ")[-1]

    doc = Document(); _setup(doc, 0.75)

    # ---- Title ----
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 0, 2)
    run(p, str(firm.get("name", "")).upper(), size=13, bold=True, color=NAVY)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 0, 10)
    run(p, "Internal Court-Preparation Packet — Confidential / Attorney Work Product", size=9, italic=True, color=SLATE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 6, 2)
    run(p, "SUMMARY EJECTMENT HEARING GUIDE & LANDLORD STATEMENT", size=15, bold=True, color=NAVY)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 0, 2)
    run(p, "%s  v.  %s" % (land.get("name", ""), tname), size=12, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 0, 10)
    run(p, "%s  •  File No. %s" % (court.get("countyCourt", ""), court.get("fileNo", "")), size=10, color=SLATE)

    # ---- Case snapshot ----
    h2(doc, "Case Snapshot")
    snap = [
        ("Hearing", [("%s at %s" % (court.get("hearingDate", ""), court.get("hearingTime", "")), {"bold": True}),
                     ("  —  %s, %s" % (court.get("courthouse", ""), court.get("room", "")), {})]),
        ("File / Case No.", "%s (%s)" % (court.get("fileNo", ""), court.get("division", ""))),
        ("Landlord (Plaintiff)", land.get("name", "")),
        ("Managing Agent", "%s (Firm License %s)" % (firm.get("name", ""), firm.get("license", ""))),
        ("Tenant (Defendant)", tname),
        ("Property", "%s (%s County)" % (prop.get("address", ""), prop.get("county", ""))),
        ("Claim", "Summary ejectment (possession) for nonpayment of rent + money judgment for rent due"),
        ("Rent charged (month at issue)", [(d["FULL"], {"bold": True}), ("  (%s)" % claim.get("rentMonthLabel", ""), {})]),
        ("Rent due through court date", [(d["PRO"], {"bold": True}),
                                          ("  (%d days × %s = %s accrued%s)" % (
                                              d["prorated_days"], d["D"], d["ACCRUED"],
                                              "; less %s paid" % money(d["current_paid"]) if d["current_paid"] else ""), {})]),
        ("Full-month alternative", [(d["MONTHASK"], {"bold": True}),
                                     ("  (%s charged%s — if the magistrate awards the whole month)" % (
                                         d["CHARGED"], (" − %s paid" % d["PAID"]) if d["current_paid"] else ""), {})]),
        ("Court costs to claim", [(d["COSTS"], {"bold": True}), ("  (%s)" % d["costs_breakdown"], {})]),
        ("Total money ask", [(d["TOTAL"], {"bold": True})]),
        ("Prepared for", "The %s licensed broker appearing for the Landlord" % firm.get("name", "")),
    ]
    fact_table(doc, snap, 2.4)

    callout(doc, "HOW TO USE THIS PACKET", [
        [("This packet has six parts: (1) the ", {}), ("Landlord's Statement", {"bold": True}),
         (" you will testify from; (2) a one-page ", {}), ("Fact Sheet", {"bold": True}),
         ("; (3) a step-by-step ", {}), ("Court Guide", {"bold": True}), ("; (4) ", {}),
         ("Anticipated Defenses & Responses", {"bold": True}), ("; (5) the ", {}),
         ("Exhibit Index & Document Checklist", {"bold": True}), ("; and (6) a ", {}),
         ("Numbers & Dates cheat card", {"bold": True}),
         (". Read Parts 1–4 the night before. Bring the whole packet plus the physical exhibits.", {})],
    ], LIGHT, RULE)

    first_lines = [
        [("This is an internal preparation aid, not legal advice. ", {"bold": True}),
         ("It was not prepared by an attorney. Have the firm's counsel review it — especially the "
          "representation-authority and \"tenant's right to pay and stay\" points — before the hearing.", {})],
        [("Two things can end this case in the tenant's favor even if everything else is perfect: ", {"bold": True}),
         ("(a) if they pay all rent in arrears plus court costs before judgment, the eviction stops by law; and "
          "(b) a money judgment requires personal service or their appearance. Both are covered inside.", {})],
    ]
    if d["has_payments"]:
        first_lines.append([("A payment has been applied on this ledger. ", {"bold": True}),
                            ("The money ask below already reflects it. Disclose the payment to the magistrate.", {})])
    if d["mismatch"]:
        first_lines.append([("LEDGER CHECK: ", {"bold": True}), (d["mismatch"], {})])
    callout(doc, "IMPORTANT — READ FIRST", first_lines, WARN, WARNB)

    # ---- Part 1 ----
    h1(doc, "Part 1 — Landlord's Statement (Opening)")
    para(doc, "Say this when the magistrate asks the landlord to present its case. Keep it short and factual.",
         size=10, italic=True, color=SLATE)
    para(doc, [('"My name is ______________________ and I am acting as the agent for the landlord of this property. '
                'The landlord is %s, and my company, %s, is the management company for the property.' % (
                    _landlord(land), firm.get("name", "")), {})])
    para(doc, [("The tenant, %s, rents %s. The rent is %s a month, due on %s. They did not pay the %s, which was due %s, "
                "and it is still unpaid. We sent a written demand for possession on %s, they did not pay, and we filed this case on %s." % (
                    tname, prop.get("address", ""), d["FULL"], rent.get("dueDayText", "the 1st"),
                    claim.get("rentMonthLabel", "rent"), _pretty(claim.get("rentDueDate")),
                    _short_date(demand.get("sentDate")), _short_date(filing.get("filedDate"))), {})])
    if d["has_payments"]:
        para(doc, [("A partial payment of %s was received and applied to the oldest outstanding charges; " % money(
            sum(a["amount"] for a in d["allocations"])), {}),
                   ("the balance of the rent remains unpaid.", {})])
    para(doc, "We are asking the court for possession of the property and a judgment for the unpaid rent plus court costs.")
    para(doc, 'I have the lease, the tenant ledger, and the written Notice and Demand for Possession here — may I enter them into evidence?"')
    extra = [
        [("Since when / lease: ", {"bold": True}), ("Tenant since %s; current written lease signed %s." % (
            tenancy.get("moveIn", "—"), tenancy.get("leaseSignedDate", "—")), {})],
        [("Proof of nonpayment: ", {"bold": True}), ('"I have the tenant ledger showing the rent charged and the balance unpaid."', {})],
    ]
    if case.get("history") or case.get("background"):
        extra.append([("History: ", {"bold": True}), (case.get("history") or case.get("background"), {})])
    extra.append([("Authority: ", {"bold": True}),
                  ("licensed broker with the managing firm, with knowledge of the records; management agreement available if requested.", {})])
    callout(doc, "If the court asks for more", extra, LIGHT, RULE)

    # ---- Part 2 ----
    h1(doc, "Part 2 — Quick Fact Sheet")
    h2(doc, "Parties & property")
    fact_table(doc, [
        ("Landlord", "%s%s" % (land.get("name", ""), (" (%s)" % land["address"]) if land.get("address") else "")),
        ("Agent / firm", "%s — Firm Lic. %s; %s, Property Manager, Lic. %s" % (
            firm.get("name", ""), firm.get("license", ""), firm.get("managerName", ""), firm.get("managerLicense", ""))),
        ("Tenant", "%s%s" % (tname, (" — %s" % ten["phone"]) if ten.get("phone") else "")),
        ("Property", "%s (%s County)%s" % (prop.get("address", ""), prop.get("county", ""),
                                            (" — %s" % prop["type"]) if prop.get("type") else "")),
        ("Move-in", tenancy.get("moveIn", "—")),
    ], 2.4)

    h2(doc, "Lease terms that matter")
    lease_rows = [
        ("Lease form / date", "%s; signed %s" % (tenancy.get("leaseForm", "—"), tenancy.get("leaseSignedDate", "—"))),
        ("Term", tenancy.get("term", "—")),
        ("Rent", "%s / month, due %s" % (d["FULL"], rent.get("dueDayText", "the 1st"))),
    ]
    if tenancy.get("para_noNotice"):
        lease_rows.append(("No notice needed", "Lease ¶%s: no notice or demand required to require payment" % tenancy["para_noNotice"]))
    if tenancy.get("para_immediateBreach"):
        lease_rows.append(("Immediate breach", "Lease ¶%s: failure to pay rent when due is an immediate breach" % tenancy["para_immediateBreach"]))
    if rent.get("lateFee"):
        lease_rows.append(("Late fee", "%s%s" % (money(rent["lateFee"]), (" (%s)" % rent["lateFeeNote"]) if rent.get("lateFeeNote") else "")))
    if tenancy.get("para_partialRent"):
        lease_rows.append(("Partial rent", "Lease ¶%s & G.S. 42-26: taking partial rent does NOT waive the eviction" % tenancy["para_partialRent"]))
    if tenancy.get("militaryAddendum"):
        lease_rows.append(("Military status", tenancy["militaryAddendum"]))
    fact_table(doc, lease_rows, 2.4)

    h2(doc, "The money (tenant ledger%s)" % ((", as of %s" % ledger["asOf"]) if ledger.get("asOf") else ""))
    lrows = []
    for r in (ledger.get("rows") or []):
        is_pay = (r.get("type") or "charge").lower() == "payment"
        amt = money(r.get("amount"))
        lrows.append([str(r.get("date", "")), str(r.get("item", "")), ("(%s)" % amt) if is_pay else amt])
    if ledger.get("balance") is not None:
        lrows.append(["", [("Ledger balance", {"bold": True})], [(money(ledger["balance"]), {"bold": True})]])
    if lrows:
        grid_table(doc, ["Date", "Item", "Amount"], lrows, [1.1, 4.5, 1.5])

    if d["has_payments"]:
        h2(doc, "How the payment was applied (oldest charge first)")
        arows = [[a["payment_date"] or "", a["applied_to"] or "", money(a["amount"])] for a in d["allocations"]]
        if d["unapplied"]:
            arows.append(["", [("Not applied to any charge — verify", {"bold": True})], [(money(d["unapplied"]), {"bold": True})]])
        grid_table(doc, ["Payment date", "Applied to", "Amount"], arows, [1.6, 4.0, 1.5])
        callout(doc, "PARTIAL PAYMENT — DISCLOSE THIS", [
            [("A payment was received and applied above. ", {"bold": True}),
             ("The rent ask has been reduced accordingly — you are asking only for rent accrued through the court "
              "date that remains unpaid.", {})],
            [("A partial payment does not stop the eviction ", {"bold": True}),
             ("(G.S. 42-26%s). Only payment of ALL rent in arrears plus court costs before judgment does (G.S. 42-33)." % (
                 "; lease ¶%s" % tenancy["para_partialRent"] if tenancy.get("para_partialRent") else ""), {})],
            [("Tell the magistrate about the payment ", {"bold": True}),
             ("before they find it on the ledger, and confirm with counsel how the firm treats partial payments.", {})],
        ], WARN, WARNB)

    ask_lines = [
        [("Ask for two things: ", {"bold": True}), ("possession of the property, and a money judgment for the rent due plus this case's ", {}),
         ("court costs", {"bold": True}), (".", {})],
        [("Rent — through the court date: ", {"bold": True}),
         ("daily rent is %s (%s ÷ %d days); %d days accrued = %s." % (
             d["D"], d["FULL"], d["days_in_month"], d["prorated_days"], d["ACCRUED"]), {})],
    ]
    if d["current_paid"]:
        ask_lines.append([("Less payments applied to that rent: ", {"bold": True}),
                          ("%s, so the rent ask is " % money(d["current_paid"]), {}), (d["PRO"], {"bold": True}), (".", {})])
    if d["prior_rent_unpaid"]:
        ask_lines.append([("Includes unpaid rent from a prior month: ", {"bold": True}), (money(d["prior_rent_unpaid"]), {}), (".", {})])
    ask_lines.append([("Court costs: ", {"bold": True}), ("%s (%s), all shown on the ledger." % (d["COSTS"], d["costs_breakdown"]), {})])
    if d["fees_outstanding"]:
        ask_lines.append([("Leave out %s of late/add-on fees. " % money(d["fees_outstanding"]), {"bold": True}),
                          ("Magistrates typically award rent and costs only. Don't muddy the ask.", {})])
    ask_lines.append([("Alternative — the whole month: ", {"bold": True}),
                      ("if the magistrate awards the full month rather than rent accrued through the court date, "
                       "the rent figure is %s (%s charged%s). Lead with %s and let the magistrate set the amount." % (
                           d["MONTHASK"], d["CHARGED"],
                           (" − %s paid" % d["PAID"]) if d["current_paid"] else "", d["PRO"]), {})])
    ask_lines.append([("TOTAL MONEY ASK: ", {"bold": True}), (d["TOTAL"], {"bold": True}),
                      ("  (or %s + %s costs on the full-month basis)" % (d["MONTHASK"], d["COSTS"]), {})])
    callout(doc, "What to ask the magistrate for", ask_lines, GREEN, GREENB)

    # ---- Part 3 ----
    h1(doc, "Part 3 — Step-by-Step Court Guide")
    h2(doc, "The day before")
    bullet(doc, "Confirm the exhibits are printed and in order (Part 5). Bring three copies of each key document: you, the magistrate, the tenant.")
    bullet(doc, [("Confirm how the tenant was served. ", {"bold": True}),
                 ("Check the sheriff's return on the summons — it decides whether a money judgment is available. "
                  "If you can't tell, call the Clerk of Superior Court's small claims office before the hearing.", {})])
    bullet(doc, "Have the management agreement handy in case anyone questions your authority to act for the Landlord.")
    bullet(doc, [("Phones are not allowed inside the courthouse. ", {"bold": True}),
                 ("Leave your cell phone in the car and print everything you need on paper.", {})])

    h2(doc, "Getting there & checking in")
    numbered(doc, [("Arrive by about %s" % court.get("arriveTime", ""), {"bold": True}),
                   (" for the %s session at %s, %s. Allow time for parking and security screening." % (
                       court.get("hearingTime", ""), court.get("courthouseShort", ""), court.get("room", "")), {})])
    numbered(doc, "Check in with the clerk/bailiff so they know the Landlord's side is present. Small claims runs as a calendar — several cases are set for the same time and called one by one.")
    numbered(doc, "If the tenant is not present when the case is called, tell the court you are ready to proceed; the magistrate can still hear the case and, if they were properly served, enter judgment.")

    h2(doc, "When your case is called")
    numbered(doc, [("State who you are: ", {"bold": True}),
                   ("your name, that you are a licensed broker with %s, the managing agent for the Landlord, and that you are here with knowledge of the records." % firm.get("name", ""), {})])
    numbered(doc, [("Present the claim ", {"bold": True}),
                   ("using the Part 1 statement: the lease, the %s rent due %s and unpaid, the written demand, and the filing date." % (
                       claim.get("rentMonthLabel", ""), _pretty(claim.get("rentDueDate"))), {})])
    numbered(doc, [("Hand up the exhibits ", {"bold": True}), ("as you mention them — lease, ledger, notice, file-stamped complaint.", {})])
    numbered(doc, [("Make the ask: ", {"bold": True}),
                   ("possession of the property and a money judgment of %s (rent %s + costs %s)." % (d["TOTAL"], d["PRO"], d["COSTS"]), {})])
    numbered(doc, [("Answer the magistrate's questions ", {"bold": True}),
                   ("directly and only from what you know. If you don't know something, say so. Stay calm and factual.", {})])
    numbered(doc, [("Listen for the ruling. ", {"bold": True}), ("The magistrate usually decides on the spot.", {})])

    h2(doc, "Proving the case — the four things the court needs")
    grid_table(doc, ["The court needs to see…", "Your proof"], [
        ["A landlord–tenant relationship / lease", "The signed %s lease (Exhibit A)" % tenancy.get("leaseForm", "lease")],
        ["Rent was due and is unpaid", "The tenant ledger showing the rent charged and the balance unpaid (Exhibit B)"],
        ["A demand was made", "The %s Notice & Demand for Possession (Exhibit C)" % _short_date(demand.get("sentDate"))],
        ["The action was properly filed & timed", "File-stamped complaint filed %s (Exhibit D)" % _short_date(filing.get("filedDate"))],
    ], [3.2, 3.9])

    h2(doc, "The money judgment — confirm service")
    svc = case.get("service", {})
    svc_line = svc.get("method") or "not stated in the documents provided — check the sheriff's return"
    callout(doc, "This determines whether you can get the money judgment today", [
        [("A North Carolina magistrate can enter a ", {}), ("money judgment", {"bold": True}), (" only if the tenant was ", {}),
         ("personally served", {"bold": True}), (" or ", {}), ("shows up at the hearing", {"bold": True}),
         (". If served only by posting on the door and they do not appear, the court can grant ", {}),
         ("possession only", {"bold": True}), (" — not the money.", {})],
        [("Service on this case: ", {"bold": True}), (svc_line, {})],
        [("Action: ", {"bold": True}),
         ("confirm from the sheriff's return before the hearing. If they appear in the courtroom, the money judgment "
          "is available regardless of how they were served — so still make the full ask.", {})],
    ], WARN, WARNB)

    h2(doc, "After the ruling — what to expect")
    bullet(doc, [("Appeal window: ", {"bold": True}), ("either side has 10 days to appeal the magistrate's decision to District Court.", {})])
    bullet(doc, [("If the 10th day falls on a weekend or holiday: ", {"bold": True}),
                 ("the deadline rolls forward to the next day the Clerk's office is open. Confirm the exact date with the Clerk before requesting the Writ of Possession.", {})])
    bullet(doc, [("No immediate lockout: ", {"bold": True}),
                 ("the writ of possession cannot issue until the appeal period passes. Never attempt a self-help lockout — removals go through the sheriff on a writ.", {})])
    bullet(doc, [("If the tenant appeals and wants to stay: ", {"bold": True}),
                 ("they must pay the rent in arrears into court and sign an undertaking to keep paying rent as it comes due.", {})])
    bullet(doc, [("Note the result: ", {"bold": True}),
                 ("write down what was awarded (possession, money, costs) and whether either side gave notice of appeal, and report to the office.", {})])

    # ---- Part 4 ----
    h1(doc, "Part 4 — Anticipated Tenant Defenses & Your Response")
    para(doc, "The tenant may raise one of these. Stay factual; don't argue — just answer.", size=10, italic=True, color=SLATE)
    partial_para = "; lease ¶%s" % tenancy["para_partialRent"] if tenancy.get("para_partialRent") else ""
    grid_table(doc, ["If the tenant says…", "Your response"], [
        ['"I paid" / "I paid part of it."',
         "Point to the ledger and the allocation table. Partial payment does not stop the eviction (G.S. 42-26%s). "
         "Only payment of ALL rent in arrears plus court costs before judgment does." % partial_para],
        ['"I want to pay now and stay."',
         "Under G.S. 42-33 they may pay all rent due plus court costs before judgment and stop the eviction. If they tender "
         "full payment, tell the magistrate you'll confirm the exact payoff (rent %s + costs %s). This is the Landlord's "
         "decision — call the office." % (d["PRO"], d["COSTS"])],
        ['"There are repair problems."',
         "Repairs are generally a separate issue and not a defense to nonpayment unless habitability was formally raised "
         "(and in some cases rent paid into court). Ask whether any written request was made; report specifics to the office."],
        ['"I never got notice."',
         "A written Notice & Demand was sent %s (Exhibit C)%s." % (
             _short_date(demand.get("sentDate")),
             "; the lease also requires no notice for rent to be due (¶%s)" % tenancy["para_noNotice"] if tenancy.get("para_noNotice") else "")],
        ['"You filed too soon."',
         "Rent was due %s and the complaint was filed %s." % (_pretty(claim.get("rentDueDate")), _short_date(filing.get("filedDate")))],
        ['"The lease is ending anyway."',
         "This case is about nonpayment of rent, not holdover. The rent is owed now and the Landlord is entitled to possession and the rent."],
        ["Asks for more time / a payment plan.",
         "Press for judgment today. Don't agree to a continuance or payment plan for a promise — they still have 10 days "
         "after judgment before any lockout to pay in full. Only full payment of rent + costs before the ruling stops the case."],
    ], [2.4, 4.7])

    callout(doc, 'THE MOST LIKELY SCENARIO — "pay and stay"', [
        [("Under ", {}), ("G.S. 42-33", {"bold": True}),
         (", if the tenant pays all rent in arrears plus the court costs of this action before the magistrate rules, the case "
          "ends and they stay. That is their legal right and there is nothing to argue. If it happens: confirm the payoff figure, "
          "get certified/guaranteed funds if possible, and let the office know. A partial payment does NOT end the case.", {})],
    ], LIGHT, RULE)
    callout(doc, "IF THE TENANT ASKS FOR MORE TIME TO PAY — THE PLAN", [
        [("Default: press for the judgment today. ", {"bold": True}),
         ("Ask for possession AND the money judgment now. Do not agree to a continuance or payment plan in exchange for a promise.", {})],
        [("Why this is still fair: ", {"bold": True}),
         ("winning today does not put them on the street. The Writ of Possession cannot issue for 10 days, so they still have "
          "that window to pay in full. Pressing now protects the Landlord's position and starts the 10-day clock.", {})],
        [("Don't freelance a deal. ", {"bold": True}),
         ("Any special arrangement is the Landlord's decision — step out and call the office before agreeing to anything.", {})],
    ], WARN, WARNB)

    # ---- Part 5 ----
    h1(doc, "Part 5 — Exhibit Index & Document Checklist")
    para(doc, "Bring the originals plus two copies of each. Check the box as you pack it.", size=10, italic=True, color=SLATE)
    ex_rows = []
    for e in (case.get("exhibits") or []):
        if isinstance(e, dict):
            ex_rows.append(["☐", e.get("label", ""), e.get("document", ""), e.get("why", "")])
        elif isinstance(e, (list, tuple)) and len(e) >= 3:
            ex_rows.append(["☐", e[0], e[1], e[2]])
    ex_rows.append(["☐", "—", "Management agreement / authority to act for Landlord", "In case your authority is questioned"])
    ex_rows.append(["☐", "—", "This packet + the quick-reference card", "Statement, fact sheet, guide"])
    grid_table(doc, ["✓", "Ex.", "Document", "Why you have it"], ex_rows, [0.4, 0.5, 3.3, 2.9])

    h2(doc, "Contacts")
    contacts = []
    if court.get("clerkLine"):
        contacts.append(("Small claims / Clerk", court["clerkLine"]))
    contacts.append(("Office", "%s — %s%s" % (firm.get("name", ""), firm.get("phone", ""),
                                              (" — %s" % firm["email"]) if firm.get("email") else "")))
    if firm.get("managerName"):
        contacts.append((firm["managerName"], "Property Manager — reach through the office for any judgment call (e.g., a pay-and-stay offer)"))
    fact_table(doc, contacts, 2.4)

    # ---- Part 6 ----
    h1(doc, "Part 6 — Numbers & Dates Cheat Card")
    para(doc, "Tear this page off and keep it on top.", size=10, italic=True, color=SLATE)
    cheat = [
        [[("Hearing", {"bold": True})], [("%s — %s — %s, %s" % (
            str(court.get("hearingDate", "")).split(", ", 1)[-1], court.get("hearingTime", ""),
            court.get("courthouseShort", ""), court.get("room", "")), {"bold": True})]],
        ["File No.", "%s (%s)" % (court.get("fileNo", ""), court.get("countyCourt", ""))],
        ["Landlord", land.get("name", "")],
        ["Tenant", tname],
        ["Property", prop.get("address", "")],
        ["Monthly rent", "%s, due %s" % (d["FULL"], rent.get("dueDayText", "the 1st"))],
        ["Daily rent", "%s/day  (%s ÷ %d)" % (d["D"], d["FULL"], d["days_in_month"])],
        ["Rent accrued through court date", "%s  (%d days)" % (d["ACCRUED"], d["prorated_days"])],
    ]
    if d["current_paid"]:
        cheat.append(["Less payments applied to rent", "(%s)" % money(d["current_paid"])])
    if d["prior_rent_unpaid"]:
        cheat.append(["Plus prior unpaid rent", money(d["prior_rent_unpaid"])])
    cheat += [
        [[("Rent to ask for", {"bold": True})], [(d["PRO"], {"bold": True})]],
        ["Full-month alternative", "%s  (%s charged%s)" % (d["MONTHASK"], d["CHARGED"],
            (" − %s paid" % d["PAID"]) if d["current_paid"] else "")],
        ["Court costs to ask for", "%s  (%s)" % (d["COSTS"], d["costs_breakdown"])],
        [[("TOTAL money ask", {"bold": True})], [(d["TOTAL"], {"bold": True})]],
    ]
    if d["fees_outstanding"]:
        cheat.append(["Fees left OUT of the ask", money(d["fees_outstanding"])])
    cheat += [
        ["Ledger balance (context)", money(ledger.get("balance")) if ledger.get("balance") is not None else "—"],
        ["Rent due date", _pretty(claim.get("rentDueDate"))],
        ["Demand notice sent", "%s%s" % (_short_date(demand.get("sentDate")),
                                         (" (pay-by %s)" % _short_date(demand.get("payByDate"))) if demand.get("payByDate") else "")],
        ["Complaint filed", _short_date(filing.get("filedDate"))],
        ["Appeal window after ruling", "10 days"],
        [[("If they ask for more time", {"bold": True})],
         [("Press for judgment now. Only full payment (rent + costs) before the ruling stops it; they still have 10 days before any lockout.", {"bold": True})]],
    ]
    grid_table(doc, ["Item", "Value"], cheat, [2.6, 4.5])

    callout(doc, "Three things not to forget", [
        [("Confirm how they were served ", {"bold": True}), ("— it controls whether you get the money judgment today.", {})],
        [("Ask for possession AND the money ", {"bold": True}),
         ("— %s rent + %s costs = %s; say both out loud." % (d["PRO"], d["COSTS"], d["TOTAL"]), {})],
        [("If they pay rent + costs in full before the ruling, ", {"bold": True}),
         ("the case ends by law — call the office; don't freelance.", {})],
    ], GREEN, GREENB)

    p = doc.add_paragraph(); sp(p, 12, 0)
    run(p, "Disclaimer: This packet was prepared as an internal preparation aid for a %s representative. It is not legal "
           "advice and was not prepared by an attorney. North Carolina law and local court practice can change and vary by "
           "county and magistrate; confirm current requirements with the firm's counsel and the %s County Clerk of Superior "
           "Court. Nothing here guarantees any particular outcome." % (firm.get("name", ""), prop.get("county", "")),
        size=8, italic=True, color="808080")

    _footer(doc.sections[0], "%s — %s — not legal advice" % (tlast, firm.get("name", "")))
    doc.save(out_path)
    return out_path
